#!/usr/bin/env python3
"""Regenerate the fixtures from a suite release folder of source workbooks and forms.

Usage: python3 scripts/generate-fixtures.py <AI_Governance_Sources_v3.9.2 folder> [recalc.py]

Writes fixtures/suite-v3.9.2-targets.json (exact headers, formula columns, controlled
lists, form fields, severity indicators, AIR-ID validation rules and versions) and, when
the LibreOffice recalc script is given, recalculates fixtures/ass02-risk-cases.json and
fixtures/ops02-closure-check-cases.json from the workbook formulas (inputs kept, outputs
read back). Needs openpyxl and python-docx; recalculation needs LibreOffice.
"""
import datetime
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
import openpyxl
from openpyxl.utils import column_index_from_string as CI
from openpyxl.utils import get_column_letter as L

SRC = Path(sys.argv[1])
RECALC = sys.argv[2] if len(sys.argv) > 2 else None
RELEASE = "v3.9.2"
ROOT = Path(__file__).resolve().parent.parent
FIX = ROOT / "fixtures"

out = {"suiteRelease": RELEASE,
       "generatedFrom": f"AI_Governance_Sources_{RELEASE} (30 September 2026) — read with openpyxl/python-docx by scripts/generate-fixtures.py; do not edit by hand",
       "sheets": {}, "forms": {}}


def dvlists(ws):
    res = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        f = dv.formula1
        if f.startswith('"'):
            vals = f.strip('"').split(",")
        else:
            sh, rng = f.split("!")
            sh = sh.strip("'")
            a, b = rng.replace("$", "").split(":")
            wsl = ws.parent[sh]
            vals = [wsl[f"{a[0]}{r}"].value for r in range(int(a[1:]), int(b[1:]) + 1)]
            vals = [v for v in vals if v is not None]
        for rng in str(dv.sqref).split():
            col = re.match(r"[A-Z]+", rng).group()
            end = re.match(r"[A-Z]+\d+:([A-Z]+)", rng)
            cols = [col] if not end else [L(i) for i in range(CI(col), CI(end.group(1)) + 1)]
            for c in cols:
                res[c] = vals
    return res


def custom_rule(ws, cell):
    for dv in ws.data_validations.dataValidation:
        if dv.type == "custom" and cell in dv.sqref:
            return dv.formula1
    return None


spec = [("ops02Monitoring", "AIG-OPS-02_AI_Post_Deployment_Monitoring_and_Review_Log_Proposed.xlsx", "Monitoring Log", 4),
        ("dec04GatePlan", "AIG-DEC-04_Gate_Log_Proposed.xlsx", "Gate plan", 3),
        ("dec04GateEvents", "AIG-DEC-04_Gate_Log_Proposed.xlsx", "Gate events", 3),
        ("dec04Conditions", "AIG-DEC-04_Gate_Log_Proposed.xlsx", "Conditions", 3),
        ("aims08Capa", "AIG-AIMS-08_AIMS_Nonconformity_and_Corrective_Action_Proposed.xlsx", "CAPA Log", 1),
        ("inv05MapChanges", "AIG-INV-05_Capabilities_and_System_Map_Proposed.xlsx", "Map changes", 3),
        ("inv04AssessmentSummary", "AIG-INV-04_AI_Register_Proposed.xlsx", "Assessment summary", 3),
        ("ass02TriageImport", "AIG-ASS-02_AI_Risk_Assessment_Worksheet_Proposed.xlsx", "Triage Import", 4)]
for key, f, s, r in spec:
    wb = openpyxl.load_workbook(SRC / f)
    ws = wb[s]
    hdr = [c.value for c in ws[r]]
    while hdr and hdr[-1] is None:
        hdr.pop()
    formula = [h for i, h in enumerate(hdr) if isinstance(ws.cell(r + 1, i + 1).value, str) and ws.cell(r + 1, i + 1).value.startswith("=")]
    lists = {hdr[CI(c) - 1]: v for c, v in dvlists(ws).items() if CI(c) - 1 < len(hdr)}
    ent = {"file": f, "sheet": s, "headerRow": r, "headers": hdr, "formulaColumns": formula, "lists": lists}
    if key == "ass02TriageImport":
        ent["formulaColumns"] = []
        ent["rows"] = [[ws.cell(i, j).value or "" for j in range(1, 6)] for i in range(5, 64)]
        ent["lists"] = {"Triage / assessment scope (B63)": dvlists(ws)["B"]}
        ra = wb["Risk Assessment"]
        rl = {}
        for dv in ra.data_validations.dataValidation:
            if dv.type == "list" and str(dv.sqref) in ("C50:C55", "C56", "C66"):
                rl[str(dv.sqref)] = dv.formula1.strip('"').split(",")
        ent["step4Lists"] = rl
        ent["step4Formulas"] = {c: ra[c].value for c in ["C42", "C58"]}
    if key == "ops02Monitoring":
        ent["formulas"] = {c: ws[f"{c}5"].value for c in ["AH", "AJ", "AK"]}
    if key == "dec04GateEvents":
        ent["formulas"] = {"N": ws["N4"].value}
    out["sheets"][key] = ent

wb = openpyxl.load_workbook(SRC / "AIG-DEC-04_Gate_Log_Proposed.xlsx")
ws = wb["Lists"]
out["sheets"]["dec04Lists"] = {"file": "AIG-DEC-04_Gate_Log_Proposed.xlsx", "sheet": "Lists",
                               "gates": [ws[f"A{i}"].value for i in range(5, 15)],
                               "lifecycleStages": [ws[f"B{i}"].value for i in range(5, 12)],
                               "outcomeMapping": [[ws[f"D{i}"].value, ws[f"E{i}"].value, ws[f"F{i}"].value] for i in range(5, 16)]}
wb = openpyxl.load_workbook(SRC / "AIG-OPS-02_AI_Post_Deployment_Monitoring_and_Review_Log_Proposed.xlsx")
ws = wb["Lists"]
out["sheets"]["ops02Lists"] = {"file": "AIG-OPS-02_AI_Post_Deployment_Monitoring_and_Review_Log_Proposed.xlsx", "sheet": "Lists",
                               "reassessmentTriggers": [ws[f"A{i}"].value for i in range(3, 11)]}

# AIR-ID validation rules (T-10): AIG-INV-04 AI Register column A and AIG-AGT-04 Agent Record column A.
out["airIdRules"] = {
    "AIG-INV-04 AI Register A4": custom_rule(openpyxl.load_workbook(SRC / "AIG-INV-04_AI_Register_Proposed.xlsx")["AI Register"], "A4"),
    "AIG-AGT-04 Agent Record A5": custom_rule(openpyxl.load_workbook(SRC / "AIG-AGT-04_Agent_Record_ASBOM_Proposed.xlsx")["Agent Record"], "A5"),
}


def rows_of(d):
    res = []
    for t in d.tables:
        for r in t.rows:
            cells = []
            for c in r.cells:
                if not cells or c._tc is not cells[-1]._tc:
                    cells.append(c)
            res.append([c.text for c in cells])
    return res


d = docx.Document(SRC / "AIG-OPS-03_AI_Incident_Report_Form.docx")
rows = rows_of(d)
lab = {r[0]: r[-1] for r in rows if len(r) == 2}
paras = [p.text for p in d.paragraphs]
desc2 = next(p for p in paras if p.startswith("Describe the incident"))
desc3 = next(p for p in paras if p.startswith("Describe the actual or potential impact"))
ticks = [t.strip() for p in paras if p.startswith("☐") and ("Residents" in p or "Service availability" in p) for t in p.split("☐") if t.strip()]
# Part A section 4 field labels in form order (two-column rows after the severity table, before Part B).
i4 = next(i for i, r in enumerate(rows) if r[0].startswith("4."))
iB = next(i for i, r in enumerate(rows) if r[0].startswith("PART B"))
sec4 = [r[0] for r in rows[i4 + 1:iB] if len(r) == 2]
partA = ["Reported by", "Role / team", "Contact email", "Date & time identified", "AI system / model name", "AIR-ID", "Affected UC-ID(s)",
         desc2, "How was it identified?", "Is the incident ongoing?", "Event classification"] + ticks + [desc3] + sec4
for x in partA:
    assert x in lab or x in paras or x in ticks, x
s8start = [i for i, r in enumerate(rows) if r[0].startswith("8.")][0]
s8 = [r[0] for r in rows[s8start + 1:] if len(r) == 2][:15]
assert s8[0] == "Personal data breach confirmed?" and s8[-1] == "Incident closed (date)", s8
sev = {r[0]: r for r in rows if len(r) == 3 and r[0] in ("Low", "Medium", "High", "Critical")}
out["forms"]["ops03PartA"] = {
    "file": "AIG-OPS-03_AI_Incident_Report_Form.docx", "section": "PART A — INITIAL REPORT, sections 1–4 (form order)", "fields": partA,
    "tickBoxes": ticks, "prompts": {k: lab[k] for k in partA if k in lab},
    "escalation": {k: v[2] for k, v in sev.items()},
    "whenItApplies": {k: v[1] for k, v in sev.items()},
    "mandatoryTriggers": next(p for p in paras if p.startswith("Mandatory triggers (Playbook §4.7.17")),
}
out["forms"]["ops03PartB8"] = {"file": "AIG-OPS-03_AI_Incident_Report_Form.docx", "section": "PART B — 8. Notification and Closure (form order)",
                               "fields": s8, "prompts": {k: lab[k] for k in s8}}
d = docx.Document(SRC / "AIG-DEC-03_Governance_Decision_Record.docx")
rows = rows_of(d)
lab = {r[0]: r[-1] for r in rows if len(r) == 2}
i1 = [i for i, r in enumerate(rows) if r[0].startswith("1.")][0]
i3 = [i for i, r in enumerate(rows) if r[0].startswith("3.")][0]
dec = [r[0] for r in rows[i1 + 1:i3] if len(r) == 2]
out["forms"]["dec03Reference"] = {"file": "AIG-DEC-03_Governance_Decision_Record.docx", "section": "1. Decision Reference and 2. Decision (form order)",
                                  "fields": dec, "prompts": {k: lab[k] for k in dec}}


def ver_xlsx(f, sheet, cell):
    return openpyxl.load_workbook(SRC / f)[sheet][cell].value


out["versions"] = {
    "AIG-OPS-02": ver_xlsx("AIG-OPS-02_AI_Post_Deployment_Monitoring_and_Review_Log_Proposed.xlsx", "Summary", "B8"),
    "AIG-DEC-04": ver_xlsx("AIG-DEC-04_Gate_Log_Proposed.xlsx", "Document Control", "B7"),
    "AIG-AIMS-08": ver_xlsx("AIG-AIMS-08_AIMS_Nonconformity_and_Corrective_Action_Proposed.xlsx", "Summary", "B8"),
    "AIG-ASS-02": ver_xlsx("AIG-ASS-02_AI_Risk_Assessment_Worksheet_Proposed.xlsx", "Document Control", "B7"),
    "AIG-INV-05": ver_xlsx("AIG-INV-05_Capabilities_and_System_Map_Proposed.xlsx", "Document Control", "B7"),
    "AIG-INV-04": ver_xlsx("AIG-INV-04_AI_Register_Proposed.xlsx", "Document Control", "B7")}
for code, f in [("AIG-OPS-03", "AIG-OPS-03_AI_Incident_Report_Form.docx"), ("AIG-DEC-03", "AIG-DEC-03_Governance_Decision_Record.docx")]:
    rows = rows_of(docx.Document(SRC / f))
    out["versions"][code] = next(r[1] for r in rows if r[0] == "Version")
(FIX / f"suite-{RELEASE}-targets.json").write_text(json.dumps(out, ensure_ascii=False, indent=1))
print("targets written:", json.dumps(out["versions"]))

if not RECALC:
    sys.exit(0)


def recalc(p):
    j = json.loads(subprocess.run(["python3", RECALC, str(p), "180"], capture_output=True, text=True).stdout)
    assert j["status"] == "success" and j["total_errors"] == 0, j


def num(v):
    return int(v) if isinstance(v, float) and v.is_integer() else v


tmp = Path(tempfile.mkdtemp())
# ---- AIG-ASS-02 risk arithmetic cases (inputs kept, outputs recalculated) ----
old = json.loads((FIX / "ass02-risk-cases.json").read_text())
f = "AIG-ASS-02_AI_Risk_Assessment_Worksheet_Proposed.xlsx"
cases = []
for i, c in enumerate(old["cases"]):
    wb = openpyxl.load_workbook(SRC / f)
    ti = wb["Triage Import"]
    for r, v in zip(range(22, 29), c["impacts"] + [c["likelihood"], c["control"]]):
        ti[f"B{r}"] = v
    p = tmp / f"ass02_{i}.xlsx"
    wb.save(p)
    recalc(p)
    ra = openpyxl.load_workbook(p, data_only=True)["Risk Assessment"]
    res = ra["C40"].value
    cases.append(dict(impacts=c["impacts"], likelihood=c["likelihood"], control=c["control"], impact=num(ra["C28"].value),
                      inherent=num(ra["C37"].value), inherentTier=ra["C38"].value,
                      residual=num(round(res, 10)) if isinstance(res, float) else res,
                      residualTier=ra["C41"].value, impactFloorFlag=ra["E28"].value or ""))
(FIX / "ass02-risk-cases.json").write_text(json.dumps({
    "source": f"{f} (suite {RELEASE}, {out['versions']['AIG-ASS-02']}): Triage Import B22–B28 set, recalculated in LibreOffice; Risk Assessment C28, C37, C38, C40, C41, E28 read back (30 September 2026).",
    "cases": cases}, ensure_ascii=False, indent=1))
print("ASS-02 cases", len(cases), "unchanged:", cases == old["cases"])

# ---- AIG-OPS-02 Monitoring Log closure check (AH), cadence (AJ), minimum sample (AK) ----
old = json.loads((FIX / "ops02-closure-check-cases.json").read_text())
f = "AIG-OPS-02_AI_Post_Deployment_Monitoring_and_Review_Log_Proposed.xlsx"
wb = openpyxl.load_workbook(SRC / f)
ws = wb["Monitoring Log"]
hdr = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(4, c).value}


def typed(v):
    if v == "":
        return None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", v):
        return datetime.datetime.strptime(v, "%Y-%m-%d")
    if re.fullmatch(r"-?\d+", v):
        return int(v)
    return v


rows_in = [c["row"] for c in old["cases"]]
for i, row in enumerate(rows_in):
    for k, v in row.items():
        ws.cell(5 + i, hdr[k]).value = typed(v)
p = tmp / "ops02.xlsx"
wb.save(p)
recalc(p)
ws = openpyxl.load_workbook(p, data_only=True)["Monitoring Log"]
blank = lambda v: "" if v is None else v
new = [{"row": row, "AH": blank(ws.cell(5 + i, CI("AH")).value), "AJ": blank(ws.cell(5 + i, CI("AJ")).value),
        "AK": blank(num(ws.cell(5 + i, CI("AK")).value))} for i, row in enumerate(rows_in)]
(FIX / "ops02-closure-check-cases.json").write_text(json.dumps({
    "source": f"{f} (suite {RELEASE}, {out['versions']['AIG-OPS-02']}), sheet Monitoring Log rows 5+; values written with openpyxl and recalculated in LibreOffice (recalc.py) on 30 September 2026. Columns AH (Closure and evidence check), AJ (Required minimum cadence), AK (Minimum sample).",
    "cases": new}, ensure_ascii=False, indent=1))
print("OPS-02 cases", len(new), "changed outputs:", sum(1 for a, b in zip(new, old["cases"]) if {k: a[k] for k in ("AH", "AJ", "AK")} != {k: b.get(k) for k in ("AH", "AJ", "AK")}))
